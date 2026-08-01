import os
import logging
import html

logger = logging.getLogger(__name__)


class ConvertError(Exception):
    pass


class ReportConverter:

    @classmethod
    def docx_to_html(cls, file_path):
        if not os.path.isfile(file_path):
            raise ConvertError(f"文件不存在: {file_path}")

        try:
            from docx import Document
        except ImportError:
            raise ConvertError("请安装 python-docx: pip install python-docx")

        doc = Document(file_path)
        html_parts = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                html_parts.append("<br>")
                continue

            escaped = html.escape(text)
            style_name = para.style.name.lower() if para.style else ""

            if "heading 1" in style_name:
                html_parts.append(f"<h1>{escaped}</h1>")
            elif "heading 2" in style_name:
                html_parts.append(f"<h2>{escaped}</h2>")
            elif "heading 3" in style_name:
                html_parts.append(f"<h3>{escaped}</h3>")
            elif "title" in style_name:
                html_parts.append(f"<h1>{escaped}</h1>")
            else:
                styled_text = cls._apply_run_styles(para)
                html_parts.append(f"<p>{styled_text}</p>")

        for table in doc.tables:
            html_parts.append(cls._table_to_html(table))

        return "\n".join(html_parts)

    @classmethod
    def _apply_run_styles(cls, paragraph):
        parts = []
        for run in paragraph.runs:
            text = html.escape(run.text)
            if not text:
                continue

            if run.bold:
                text = f"<strong>{text}</strong>"
            if run.italic:
                text = f"<em>{text}</em>"
            if run.underline:
                text = f"<u>{text}</u>"

            parts.append(text)

        return "".join(parts) if parts else html.escape(paragraph.text)

    @classmethod
    def _table_to_html(cls, table):
        rows = []
        for i, row in enumerate(table.rows):
            cells = []
            tag = "th" if i == 0 else "td"
            for cell in row.cells:
                text = html.escape(cell.text.strip())
                cells.append(f"<{tag}>{text}</{tag}>")
            rows.append("<tr>" + "".join(cells) + "</tr>")

        return (
            '<table border="1" cellpadding="5" cellspacing="0" '
            'style="border-collapse:collapse;width:100%">'
            + "".join(rows)
            + "</table>"
        )

    @classmethod
    def to_plain_text(cls, file_path):
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".docx":
            return cls._docx_to_text(file_path)
        elif ext == ".pdf":
            return cls._pdf_to_text(file_path)
        else:
            raise ConvertError(f"不支持的文件类型: {ext}")

    @classmethod
    def _docx_to_text(cls, file_path):
        try:
            from docx import Document
        except ImportError:
            raise ConvertError("请安装 python-docx")

        doc = Document(file_path)
        lines = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                lines.append(row_text)

        return "\n".join(lines)

    @classmethod
    def _pdf_to_text(cls, file_path):
        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                from PyPDF2 import PdfReader
            except ImportError:
                raise ConvertError("请安装 pypdf: pip install pypdf")

        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text.strip())

        return "\n".join(text_parts)

    @classmethod
    def get_file_info(cls, file_path):
        if not os.path.isfile(file_path):
            raise ConvertError(f"文件不存在: {file_path}")

        stat = os.stat(file_path)
        return {
            "filename": os.path.basename(file_path),
            "file_type": os.path.splitext(file_path)[1].lstrip("."),
            "file_size": stat.st_size,
            "file_size_display": cls._format_size(stat.st_size),
        }

    @staticmethod
    def _format_size(size_bytes):
        if size_bytes < 1024:
            return f"{size_bytes} B"
        elif size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} KB"
        else:
            return f"{size_bytes / 1024 / 1024:.2f} MB"
